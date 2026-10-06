import os
import re
import datetime
import openpyxl
import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "balajiinfra_system.settings.dev")
django.setup()

from payroll_management.models import Company, Employee_detail, AttendanceRecord

def parse_date_str(val):
    if not val:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    s = str(val).strip()
    for fmt in ("%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()

def clean_num(val, default=0.0):
    if val is None or val == "":
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default

def run():
    excel_path = "d:/chronoanalytics/Balaji-InfraTech-LLP/Jun Wages Sheet.xlsx"
    print(f"Loading {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["BILLING & PAY SHEET"]

    company, _ = Company.objects.get_or_create(
        id=1,
        defaults={
            "name": "Nandurbar",
            "pf": 13.0,
            "esic": 3.25,
            "lwf": 75.0,
            "uniform_charge": 300.0,
            "admin_charge": 7.0,
        }
    )

    # Clear existing to ensure clean sequence
    AttendanceRecord.objects.filter(company=company).delete()
    Employee_detail.objects.all().delete()
    print("Cleaned existing employee & attendance records.")

    emp_created = 0
    att_created_jun = 0
    att_created_sep = 0

    for r in range(9, 118):
        sno_val = ws[r][0].value
        if sno_val is None:
            continue
        try:
            sno = int(sno_val)
        except ValueError:
            continue

        emp_code = clean_str(ws[r][2].value)
        name = clean_str(ws[r][3].value)
        father_name = clean_str(ws[r][4].value)
        sex = clean_str(ws[r][5].value) or "Male"
        address = clean_str(ws[r][6].value)
        dob = parse_date_str(ws[r][7].value)
        doj = parse_date_str(ws[r][8].value)
        adhar_no = clean_str(ws[r][9].value)
        bank_name = clean_str(ws[r][10].value)
        account_no = clean_str(ws[r][11].value)
        ifsc_code = clean_str(ws[r][12].value)
        uan_no = clean_str(ws[r][13].value)
        esic_no = clean_str(ws[r][14].value)
        mobile_no = clean_str(ws[r][15].value)
        dept = clean_str(ws[r][17].value) or "MAIN MOULD ONE FORTY SEVEN"
        desig = clean_str(ws[r][18].value) or "ITI"
        category = clean_str(ws[r][19].value) or "SKILLED"
        calc_days = int(clean_num(ws[r][20].value, 27))
        work_days = clean_num(ws[r][21].value, 27)
        paid_days = clean_num(ws[r][24].value, 27)

        # Wages structure
        basic = clean_num(ws[r][46].value, 15525.0)
        hra = clean_num(ws[r][47].value, 776.25)
        cca = clean_num(ws[r][48].value, 0.0)
        incentive = clean_num(ws[r][49].value, 0.0)
        bonus = clean_num(ws[r][50].value, 1293.23)
        leave_enc = clean_num(ws[r][51].value, 776.25)
        transport = clean_num(ws[r][52].value, 1500.0)
        gross_salary = clean_num(ws[r][55].value, 19870.73)

        pf_13 = clean_num(ws[r][65].value, 1950.0)
        esic_3_25 = clean_num(ws[r][66].value, 530.0)
        lwf_comp = clean_num(ws[r][67].value, 75.0)
        uniform = clean_num(ws[r][68].value, 300.0)
        services_7 = clean_num(ws[r][69].value, 1486.0)
        ctc = clean_num(ws[r][70].value, 24211.73)

        # Deductions
        pf_12 = clean_num(ws[r][56].value, 1800.0)
        esic_0_75 = clean_num(ws[r][57].value, 123.0)
        ptax = clean_num(ws[r][58].value, 200.0)
        canteen = clean_num(ws[r][61].value, 405.0)
        lwf_emp = clean_num(ws[r][62].value, 25.0)
        total_ded = clean_num(ws[r][63].value, 2553.0)
        net_sal = clean_num(ws[r][64].value, 17318.0)

        # Create Employee in DB
        emp = Employee_detail(
            id=sno,
            balaji_emp_code=emp_code,
            salutation="Mr.",
            name=name,
            father_or_husband_name=father_name,
            sex=sex,
            marital_status="Not Married",
            current_address_line1=address,
            dob=dob,
            doj=doj,
            adhar_no=adhar_no,
            bank_name=bank_name,
            account_no=account_no,
            ifsc_code=ifsc_code,
            uan_no=uan_no,
            esic_no=esic_no,
            mobile_no=mobile_no,
            department=dept,
            designation=desig,
            category=category,
            company=company,
            calculation_days=calc_days,
            basic=basic,
            hra=hra,
            cca_site_allowance=cca,
            attendance_bonus=incentive,
            bonus=bonus,
            leave_encashment=leave_enc,
            gross_salary=gross_salary,
            transport=transport,
            pf_13=pf_13,
            esic_3_25=esic_3_25,
            uniform_shoes=uniform,
            services_7=services_7,
            ctc=ctc,
        )
        emp.save()
        emp_created += 1

        # Create Attendance & Deduction Record for June 2026
        AttendanceRecord.objects.create(
            company=company,
            attendance_date="2026-06",
            sr_no=sno,
            emp_code=emp_code,
            emp_name=name,
            mandays=work_days,
            payable_mandays=paid_days,
            ot=0.0,
            canteen=canteen,
            pf=pf_12,
            esic=esic_0_75,
            ptax=ptax,
            lwf=lwf_emp,
            total_deduction=total_ded,
        )
        att_created_jun += 1

        # Create Attendance & Deduction Record for September 2026 (current month)
        AttendanceRecord.objects.create(
            company=company,
            attendance_date="2026-09",
            sr_no=sno,
            emp_code=emp_code,
            emp_name=name,
            mandays=work_days,
            payable_mandays=paid_days,
            ot=0.0,
            canteen=canteen,
            pf=pf_12,
            esic=esic_0_75,
            ptax=ptax,
            lwf=lwf_emp,
            total_deduction=total_ded,
        )
        att_created_sep += 1

    print(f"Successfully seeded: {emp_created} employees, {att_created_jun} June attendance records, {att_created_sep} September attendance records.")

if __name__ == "__main__":
    run()
